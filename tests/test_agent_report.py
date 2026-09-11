"""Tests for dia.agent.report — ReportUploader."""

from datetime import date
from io import BytesIO

import boto3
import pytest
from docx import Document
from moto import mock_aws

from dia.agent.models import AgentResponse
from dia.agent.report import ReportUploader

BUCKET = "test-agent-reports-bucket"


@pytest.fixture
def s3_client():
    with mock_aws():
        client = boto3.client("s3", region_name="eu-west-2")
        client.create_bucket(
            Bucket=BUCKET,
            CreateBucketConfiguration={"LocationConstraint": "eu-west-2"},
        )
        yield client


def _response(**overrides) -> AgentResponse:
    fields = {
        "id": "a1b2c3d4-e5f6-4789-a123-000000000000",
        "run_date": date(2026, 9, 11),
        "department": "Home Office",
        "query": "What's the DBR for Home Office?",
        "output": "# DBR\n\nSome findings.",
    }
    fields.update(overrides)
    return AgentResponse(**fields)


# --- _slugify ---


def test_slugify_lowercases_and_hyphenates():
    assert ReportUploader._slugify("Home Office") == "home-office"


def test_slugify_strips_punctuation():
    assert ReportUploader._slugify("Department for Education & Skills!") == "department-for-education-skills"


def test_slugify_none_defaults_to_cross_government():
    assert ReportUploader._slugify(None) == "cross-government"


def test_slugify_empty_string_defaults_to_cross_government():
    assert ReportUploader._slugify("") == "cross-government"


# --- _build_key ---


def test_build_key_uses_department_slug_date_and_short_id():
    response = _response()

    key = ReportUploader._build_key(response, "docx")

    assert key == "home-office/2026-09-11_a1b2c3d4.docx"


def test_build_key_uses_cross_government_when_no_department():
    response = _response(department=None)

    key = ReportUploader._build_key(response, "md")

    assert key == "cross-government/2026-09-11_a1b2c3d4.md"


# --- _to_docx: headings, lists, paragraphs ---


def _render(markdown: str) -> Document:
    return Document(BytesIO(ReportUploader._to_docx(markdown)))


def test_to_docx_renders_headings_with_correct_levels():
    doc = _render("# Title\n\n## Subtitle\n\n### Sub-subtitle")

    styles = [p.style.name for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]

    assert styles == ["Heading 1", "Heading 2", "Heading 3"]
    assert texts == ["Title", "Subtitle", "Sub-subtitle"]


def test_to_docx_renders_bullet_list():
    doc = _render("- First\n- Second\n* Third")

    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert [p.text for p in bullets] == ["First", "Second", "Third"]


def test_to_docx_renders_numbered_list():
    doc = _render("1. First\n2. Second")

    numbered = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert [p.text for p in numbered] == ["First", "Second"]


def test_to_docx_renders_plain_paragraph():
    doc = _render("Just a plain sentence.")

    normal = [p for p in doc.paragraphs if p.style.name == "Normal"]
    assert [p.text for p in normal] == ["Just a plain sentence."]


def test_to_docx_skips_blank_lines():
    doc = _render("First paragraph.\n\n\n\nSecond paragraph.")

    texts = [p.text for p in doc.paragraphs]
    assert texts == ["First paragraph.", "Second paragraph."]


# --- _to_docx: inline formatting ---


def test_to_docx_bold_span_becomes_bold_run():
    doc = _render("**Emerald Programme** is at risk.")

    runs = doc.paragraphs[0].runs
    assert runs[0].text == "Emerald Programme"
    assert runs[0].bold is True


def test_to_docx_italic_span_becomes_italic_run():
    doc = _render("This is *important* context.")

    italic_runs = [r for r in doc.paragraphs[0].runs if r.italic]
    assert [r.text for r in italic_runs] == ["important"]


def test_to_docx_code_span_uses_monospace_font():
    doc = _render("Run `get_table_schema` first.")

    code_runs = [r for r in doc.paragraphs[0].runs if r.font.name == "Consolas"]
    assert [r.text for r in code_runs] == ["get_table_schema"]


def test_to_docx_preserves_pound_sign():
    doc = _render("Total spend: £412m across all programmes.")

    assert "£412m" in doc.paragraphs[0].text


# --- _to_docx: tables ---


def test_to_docx_renders_gfm_table_as_real_table():
    markdown = "| Supplier | Risk |\n|---|---|\n| Fujitsu | High |\n| Kainos | Low |"

    doc = _render(markdown)

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert [c.text for c in table.rows[0].cells] == ["Supplier", "Risk"]
    assert [c.text for c in table.rows[1].cells] == ["Fujitsu", "High"]
    assert [c.text for c in table.rows[2].cells] == ["Kainos", "Low"]


def test_to_docx_table_header_row_is_bold():
    markdown = "| Supplier | Risk |\n|---|---|\n| Fujitsu | High |"

    doc = _render(markdown)

    header_cell = doc.tables[0].rows[0].cells[0]
    assert header_cell.paragraphs[0].runs[0].bold is True


def test_to_docx_uses_table_grid_style():
    markdown = "| A | B |\n|---|---|\n| 1 | 2 |"

    doc = _render(markdown)

    assert doc.tables[0].style.name == "Table Grid"


def test_to_docx_does_not_treat_inline_pipes_as_table():
    doc = _render("Total: £1m | RDEL: £2m | CDEL: £3m")

    assert doc.tables == []
    assert "£1m | RDEL: £2m | CDEL: £3m" in doc.paragraphs[0].text


def test_to_docx_handles_ragged_table_rows():
    markdown = "| A | B | C |\n|---|---|---|\n| 1 | 2 |"

    doc = _render(markdown)

    table = doc.tables[0]
    assert [c.text for c in table.rows[1].cells] == ["1", "2", ""]


# --- upload() ---


def test_upload_writes_markdown_object(s3_client):
    uploader = ReportUploader(bucket=BUCKET, s3_client=s3_client)
    response = _response()

    uploader.upload(response)

    obj = s3_client.get_object(Bucket=BUCKET, Key="home-office/2026-09-11_a1b2c3d4.md")
    assert obj["Body"].read().decode() == response.output
    assert obj["ContentType"] == "text/markdown"


def test_upload_writes_docx_object(s3_client):
    uploader = ReportUploader(bucket=BUCKET, s3_client=s3_client)
    response = _response()

    uploader.upload(response)

    obj = s3_client.get_object(Bucket=BUCKET, Key="home-office/2026-09-11_a1b2c3d4.docx")
    assert obj["ContentType"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    # Confirm it's a real, openable docx.
    doc = Document(BytesIO(obj["Body"].read()))
    assert doc.paragraphs[0].text == "DBR"


def test_upload_returns_uris_and_download_urls(s3_client):
    uploader = ReportUploader(bucket=BUCKET, s3_client=s3_client)
    response = _response()

    result = uploader.upload(response)

    assert result["markdown_uri"] == f"s3://{BUCKET}/home-office/2026-09-11_a1b2c3d4.md"
    assert result["docx_uri"] == f"s3://{BUCKET}/home-office/2026-09-11_a1b2c3d4.docx"
    assert result["markdown_download_url"].startswith(f"https://{BUCKET}.s3.amazonaws.com/")
    assert result["docx_download_url"].startswith(f"https://{BUCKET}.s3.amazonaws.com/")


def test_upload_uses_cross_government_folder_when_no_department(s3_client):
    uploader = ReportUploader(bucket=BUCKET, s3_client=s3_client)
    response = _response(department=None)

    result = uploader.upload(response)

    assert result["markdown_uri"] == f"s3://{BUCKET}/cross-government/2026-09-11_a1b2c3d4.md"


def test_upload_respects_custom_expiry(s3_client):
    uploader = ReportUploader(bucket=BUCKET, s3_client=s3_client)
    response = _response()

    result = uploader.upload(response, url_expiry_seconds=3_600)

    assert "X-Amz-Expires=3600" in result["docx_download_url"]
